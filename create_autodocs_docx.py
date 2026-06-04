import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import glob

def create_document():
    doc = docx.Document()
    
    # Set default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(14)
    
    # Add Chapter Title
    heading = doc.add_heading('ГЛАВА 3. ОФОРМЛЕНИЕ РЕЗУЛЬТАТОВ И ДОСТИЖЕНИЙ ПРОЕКТА', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Intro
    p = doc.add_paragraph('Заключительный этап разработки системы электронного документооборота autodocs связан с оформлением результатов, проверкой соответствия итогового продукта первоначальным требованиям, а также проработкой взаимодействия конечных пользователей с системой. В данной главе рассматриваются ключевые аспекты, связанные с дизайном пользовательского интерфейса, организацией навигации, обеспечением кроссплатформенной адаптивности и эргономики разрабатываемой платформы. Успешность внедрения информационной системы во многом определяется тем, насколько эффективно, комфортно и интуитивно понятно пользователи могут решать свои задачи в рамках предоставленного функционала.')
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    media_files = sorted(glob.glob('/Users/rustamurinboyev/.gemini/antigravity/brain/5af44db9-5ef7-4843-a476-c19b47d71750/media*.png'))
    
    # Section 3.1
    h2 = doc.add_heading('3.1 Главная страница', level=2)
    p = doc.add_paragraph('Визуальное оформление главной страницы играет критически важную роль в формировании первого впечатления о системе autodocs. Интерфейс спроектирован с учетом современных тенденций, ориентированных на минимализм, чистоту и консистентность.')
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    if len(media_files) > 0:
        doc.add_picture(media_files[0], width=Inches(6.0))
        doc.add_paragraph('Рисунок 3.1 - Главная страница системы', style='Caption')
    
    # Section 3.2
    h2 = doc.add_heading('3.2 Панель обычного пользователя', level=2)
    p = doc.add_paragraph('Для обычных пользователей системы реализован упрощенный интерфейс, предоставляющий доступ к документам и настройкам. Логичная и интуитивно понятная навигация позволяет быстро находить необходимые функции.')
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    if len(media_files) > 1:
        doc.add_picture(media_files[1], width=Inches(6.0))
        doc.add_paragraph('Рисунок 3.2 - Панель обычного пользователя', style='Caption')

    # Section 3.3
    h2 = doc.add_heading('3.3 Панель администратора', level=2)
    p = doc.add_paragraph('Панель администратора обладает расширенными возможностями для управления системой, включая работу со всеми документами, пользователями и системными настройками. Ниже представлены скриншоты интерфейса панели администратора.')
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    if len(media_files) > 2:
        doc.add_picture(media_files[2], width=Inches(6.0))
        doc.add_paragraph('Рисунок 3.3 - Панель администратора (часть 1)', style='Caption')

    if len(media_files) > 3:
        doc.add_picture(media_files[3], width=Inches(6.0))
        doc.add_paragraph('Рисунок 3.4 - Панель администратора (часть 2)', style='Caption')

    p = doc.add_paragraph('Навигация в панели администратора организована с помощью бокового меню, что позволяет легко переключаться между разделами. Кроме того, реализован "умный поиск", использующий семантические технологии для точного нахождения необходимых документов.')
    p.paragraph_format.first_line_indent = Inches(0.5)

    # Section 3.4
    h2 = doc.add_heading('3.4 Адаптивность и эргономика', level=2)
    p = doc.add_paragraph('Интерфейс системы полностью адаптивен и корректно отображается как на настольных компьютерах, так и на мобильных устройствах. Использование современных веб-технологий позволяет достичь высокой производительности и плавности работы.')
    p.paragraph_format.first_line_indent = Inches(0.5)

    p = doc.add_paragraph('Выводы по главе:')
    p.paragraph_format.first_line_indent = Inches(0.5)
    p.runs[0].bold = True

    p = doc.add_paragraph('Разработанный пользовательский интерфейс системы autodocs отвечает современным требованиям веб-дизайна, эргономики и юзабилити. Применение стандартизированной дизайн-системы, продуманной навигации и адаптивной верстки обеспечивает высокий уровень комфорта для пользователей.')
    p.paragraph_format.first_line_indent = Inches(0.5)
    
    doc.save('/Users/rustamurinboyev/Documents/Проект/Глава_3_autodocs.docx')

if __name__ == "__main__":
    create_document()
